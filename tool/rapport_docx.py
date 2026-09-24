#!/usr/bin/env python3
"""Engendre le rapport d'activité au format Word, sur le gabarit D-CLIC.

Le gabarit de la formation (« ACTIVITE 1 DCLIC.pdf ») impose une couverture
encadrée de violet, portant le logo de l'Organisation internationale de la
Francophonie, le niveau, la session, le numéro d'activité, le titre, l'auteur
et la date, chacun dans son propre cadre de couleur. Viennent ensuite une
page laissée blanche, une page d'objectifs et une synthèse.

Word n'expose ni bordure de page ni fond de paragraphe à travers python-docx :
ces attributs sont posés directement dans le XML du document, d'où les
fonctions `_bordure_*` et `_fond` ci-dessous.

Usage : python3 tool/rapport_docx.py
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJET = Path(__file__).resolve().parent.parent
SORTIE = PROJET / "Rapport_Projet_Mekano_Afrika_GNAZOU_BERNARD.docx"
LOGO = PROJET / "documents" / "logo_dclic.png"
CAPTURES = PROJET / "captures"

AUTEUR = "GNAZOU GOUDI BERNARD"
NIVEAU = "DEVELOPPEMENT MOBILE NIVEAU APPROFONDI"
SESSION = "SESSION AOUT-2026"
ACTIVITE = "ACTIVITE : PROJET FINAL - DOCUMENTATION, TESTS, PRESENTATION, DEPOT"
TITRE = "CONCEPTION ET RÉALISATION\nDE L'APPLICATION MOBILE\nMEKANO AFRIKA"
DATE = "24/09/2026"

# Couleurs du gabarit, relevées sur la couverture d'origine.
VIOLET = "7030A0"
ORANGE = "FFC000"
ROUGE = "C00000"
BLEU = "00B0F0"
VERT_PALE = "E2EFDA"
NOIR = "000000"


# Le schéma OOXML impose l'ordre des enfants de `w:pPr` et de `w:sectPr` :
# un élément ajouté à la fin est au mieux ignoré, au pire rend le document
# invalide. D'où ces listes de successeurs, qui disent où insérer.
_APRES_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)

_APRES_SHD = _APRES_PBDR[1:]

_APRES_PGBORDERS = (
    "w:lnNumType", "w:pgNumType", "w:cols", "w:formProt", "w:vAlign",
    "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter",
    "w:docGrid", "w:printerSettings", "w:sectPrChange",
)


def _element(nom, **attributs):
    balise = OxmlElement(nom)
    for cle, valeur in attributs.items():
        balise.set(qn(f"w:{cle}"), valeur)
    return balise


def bordure_de_page(section, couleur=VIOLET, epaisseur=18):
    """Encadre la page entière, comme la couverture du gabarit."""
    bordures = _element("w:pgBorders", offsetFrom="page")
    for cote in ("top", "left", "bottom", "right"):
        bordures.append(
            _element(
                f"w:{cote}",
                val="single",
                sz=str(epaisseur),
                space="24",
                color=couleur,
            )
        )
    section._sectPr.insert_element_before(bordures, *_APRES_PGBORDERS)


def encadrer(paragraphe, couleur=NOIR, epaisseur=8):
    """Pose un filet autour d'un paragraphe."""
    pPr = paragraphe._p.get_or_add_pPr()
    bordures = _element("w:pBdr")
    for cote in ("top", "left", "bottom", "right"):
        bordures.append(
            _element(
                f"w:{cote}",
                val="single",
                sz=str(epaisseur),
                space="6",
                color=couleur,
            )
        )
    pPr.insert_element_before(bordures, *_APRES_PBDR)


def fond(paragraphe, couleur):
    """Applique un aplat de couleur derrière un paragraphe."""
    pPr = paragraphe._p.get_or_add_pPr()
    pPr.insert_element_before(
        _element("w:shd", val="clear", color="auto", fill=couleur),
        *_APRES_SHD,
    )


def ligne(document, texte="", taille=11, gras=False, italique=False,
          couleur=None, alignement=WD_ALIGN_PARAGRAPH.LEFT, espace_avant=0,
          espace_apres=6, police=None):
    """Paragraphe d'une seule fonte, réglé d'un seul appel."""
    p = document.add_paragraph()
    p.alignment = alignement
    p.paragraph_format.space_before = Pt(espace_avant)
    p.paragraph_format.space_after = Pt(espace_apres)
    r = p.add_run(texte)
    r.font.size = Pt(taille)
    r.bold = gras
    r.italic = italique
    if police:
        r.font.name = police
    if couleur:
        r.font.color.rgb = RGBColor.from_string(couleur)
    return p


def titre(document, texte, niveau=1):
    """Intertitre, dans les tons du gabarit."""
    tailles = {1: 16, 2: 13, 3: 12}
    couleurs = {1: ROUGE, 2: ROUGE, 3: "1F1F23"}
    p = ligne(
        document,
        texte,
        taille=tailles[niveau],
        gras=True,
        couleur=couleurs[niveau],
        espace_avant=18 if niveau == 1 else 12,
        espace_apres=8,
    )
    return p


def puces(document, elements, numerotees=False):
    style = "List Number" if numerotees else "List Bullet"
    for element in elements:
        p = document.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(4)
        if isinstance(element, tuple):
            tete, suite = element
            r = p.add_run(tete)
            r.bold = True
            r.font.size = Pt(11)
            r = p.add_run(suite)
            r.font.size = Pt(11)
        else:
            r = p.add_run(element)
            r.font.size = Pt(11)


def tableau(document, entetes, lignes):
    t = document.add_table(rows=1, cols=len(entetes))
    t.style = "Table Grid"
    for cellule, texte in zip(t.rows[0].cells, entetes):
        cellule.text = ""
        p = cellule.paragraphs[0]
        r = p.add_run(texte)
        r.bold = True
        r.font.size = Pt(10)
        fond(p, "F2F2F2")
    for contenu in lignes:
        cellules = t.add_row().cells
        for cellule, texte in zip(cellules, contenu):
            cellule.text = ""
            r = cellule.paragraphs[0].add_run(texte)
            r.font.size = Pt(10)
    document.add_paragraph()
    return t


def capture(document, fichier, legende, largeur=Cm(7.2)):
    """Insère une capture d'écran, légendée."""
    chemin = CAPTURES / fichier
    if not chemin.exists():
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(chemin), width=largeur)
    ligne(
        document,
        legende,
        taille=9,
        italique=True,
        couleur="595959",
        alignement=WD_ALIGN_PARAGRAPH.CENTER,
        espace_apres=12,
    )


def couverture(document):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    for marge in ("left_margin", "right_margin"):
        setattr(section, marge, Cm(2.2))
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    bordure_de_page(section)

    # Logo à gauche, niveau à droite : un tableau sans filet tient les deux.
    entete = document.add_table(rows=1, cols=2)
    entete.autofit = False
    entete.columns[0].width = Cm(8.2)
    entete.columns[1].width = Cm(8.2)

    cellule_logo = entete.rows[0].cells[0]
    cellule_logo.text = ""
    if LOGO.exists():
        cellule_logo.paragraphs[0].add_run().add_picture(
            str(LOGO), width=Cm(6.4)
        )

    cellule_niveau = entete.rows[0].cells[1]
    cellule_niveau.text = ""
    p = cellule_niveau.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(NIVEAU)
    r.font.size = Pt(11)
    encadrer(p, ORANGE, 8)

    p = cellule_niveau.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run(SESSION)
    r.font.size = Pt(10)
    r.bold = True
    r.italic = True

    for _ in range(4):
        ligne(document, espace_apres=0)

    p = ligne(document, ACTIVITE, taille=13, espace_apres=0)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.6)
    encadrer(p, ROUGE, 8)

    ligne(document, espace_apres=0)
    ligne(document, espace_apres=0)

    # Le titre, dans le grand cadre bleu.
    premier = None
    for index, texte in enumerate(TITRE.split("\n")):
        p = ligne(
            document,
            texte,
            taille=22,
            gras=True,
            alignement=WD_ALIGN_PARAGRAPH.CENTER,
            espace_avant=14 if index == 0 else 0,
            espace_apres=14 if index == len(TITRE.split("\n")) - 1 else 4,
        )
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.right_indent = Cm(0.4)
        if premier is None:
            premier = p
        # Le cadre englobe les trois lignes : seules les bordures haute et
        # basse changent d'une ligne à l'autre.
        pPr = p._p.get_or_add_pPr()
        bordures = _element("w:pBdr")
        cotes = {"left": True, "right": True,
                 "top": index == 0,
                 "bottom": index == len(TITRE.split("\n")) - 1}
        for cote, actif in cotes.items():
            bordures.append(
                _element(
                    f"w:{cote}",
                    val="single" if actif else "nil",
                    sz="18" if actif else "0",
                    space="8",
                    color=BLEU if actif else "auto",
                )
            )
        pPr.insert_element_before(bordures, *_APRES_PBDR)

    for _ in range(3):
        ligne(document, espace_apres=0)

    p = ligne(document, "RÉALISÉ PAR :", taille=11, espace_apres=0)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(11.0)
    encadrer(p, NOIR, 6)

    p = ligne(document, AUTEUR, taille=11, espace_avant=2, espace_apres=0)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(9.0)
    fond(p, VERT_PALE)

    for _ in range(5):
        ligne(document, espace_apres=0)

    p = ligne(
        document,
        f"DATE : {DATE}",
        taille=12,
        alignement=WD_ALIGN_PARAGRAPH.CENTER,
        espace_apres=0,
    )
    p.paragraph_format.left_indent = Cm(10.5)
    encadrer(p, NOIR, 6)


def page_blanche(document):
    """Le gabarit laisse la deuxième page vierge."""
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    ligne(document, espace_apres=0)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def main() -> int:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)

    couverture(document)
    page_blanche(document)
    from contenu_rapport import objectifs, corps  # noqa: E402

    objectifs(document, ligne, titre, puces, capture)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    corps(document, ligne, titre, puces, tableau, capture)

    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    document.save(SORTIE)
    print(f"{SORTIE.relative_to(PROJET)}, {SORTIE.stat().st_size // 1024} ko")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
